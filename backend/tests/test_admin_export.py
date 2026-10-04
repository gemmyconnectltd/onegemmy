import csv
import io

from openpyxl import load_workbook

from app.modules.admin.export import (
    COLUMN_DEFS,
    DEFAULT_COLUMNS,
    build_users_export,
    resolve_columns,
)

SAMPLE_ROWS = [
    {
        "id": "1", "email": "robert@example.com", "full_name": "Robert Niyitanga", "phone": "+250788000000",
        "role": "admin", "is_active": True, "is_superuser": False,
        "created_at": "2026-01-15T10:30:00+00:00", "updated_at": "2026-02-01T09:00:00+00:00",
        "tenant_id": "t1", "tenant_name": "FreshMart", "branch_id": "b1", "branch_name": "Main Branch",
    },
    {
        "id": "2", "email": "jane@example.com", "full_name": "Jane Doe", "phone": None,
        "role": "member", "is_active": False, "is_superuser": False,
        "created_at": "2026-03-20T14:00:00+00:00", "updated_at": None,
        "tenant_id": None, "tenant_name": None, "branch_id": None, "branch_name": None,
    },
]


# ── Security: never expose auth-sensitive fields ──────────────────────────────

def test_column_defs_never_expose_sensitive_fields():
    forbidden = {"password", "hashed_password", "token", "refresh_token", "access_token", "otp", "secret", "session"}
    for key in COLUMN_DEFS:
        assert not any(bad in key.lower() for bad in forbidden), f"column '{key}' looks sensitive"


def test_resolve_columns_ignores_unknown_keys():
    """A client can't smuggle an arbitrary field through `columns` — only
    keys present in COLUMN_DEFS ever make it through."""
    resolved = resolve_columns(["full_name", "hashed_password", "email", "nonsense"])
    assert resolved == ["full_name", "email"]


def test_resolve_columns_defaults_when_none_requested():
    assert resolve_columns(None) == DEFAULT_COLUMNS
    assert resolve_columns([]) == DEFAULT_COLUMNS


def test_resolve_columns_falls_back_to_default_when_all_invalid():
    assert resolve_columns(["password", "nonsense"]) == DEFAULT_COLUMNS


# ── CSV ─────────────────────────────────────────────────────────────────────

def test_csv_export_contains_headers_and_rows():
    filename, content, media_type = build_users_export(SAMPLE_ROWS, DEFAULT_COLUMNS, "csv", "admin@pesaa.io", {})
    assert filename.endswith(".csv")
    assert media_type.startswith("text/csv")

    text = content.decode("utf-8-sig")
    reader = list(csv.reader(io.StringIO(text)))
    header = reader[0]
    assert header == [COLUMN_DEFS[c][0] for c in DEFAULT_COLUMNS]
    assert len(reader) == 1 + len(SAMPLE_ROWS)
    assert "Robert Niyitanga" in reader[1]
    assert "FreshMart" in reader[1]


def test_csv_filename_distinguishes_filtered_vs_all():
    unfiltered, _, _ = build_users_export(SAMPLE_ROWS, DEFAULT_COLUMNS, "csv", "a@pesaa.io", {})
    filtered, _, _ = build_users_export(SAMPLE_ROWS, DEFAULT_COLUMNS, "csv", "a@pesaa.io", {"search": "Robert"})
    assert "filtered" not in unfiltered
    assert "filtered" in filtered


def test_csv_never_contains_password_column_even_if_requested():
    _, content, _ = build_users_export(SAMPLE_ROWS, resolve_columns(["hashed_password"]), "csv", "a@pesaa.io", {})
    assert "password" not in content.decode("utf-8-sig").lower()


# ── XLSX ────────────────────────────────────────────────────────────────────

def test_xlsx_export_has_bold_header_and_correct_rows():
    _, content, media_type = build_users_export(SAMPLE_ROWS, DEFAULT_COLUMNS, "xlsx", "admin@pesaa.io", {"status": "active"})
    assert media_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    wb = load_workbook(io.BytesIO(content))
    ws = wb.active

    # Find the header row (after the report's title/meta block)
    header_row = next(r for r in range(1, 20) if ws.cell(r, 1).value == COLUMN_DEFS[DEFAULT_COLUMNS[0]][0])
    assert ws.cell(header_row, 1).font.bold is True
    assert ws.freeze_panes == f"A{header_row + 1}"
    assert ws.auto_filter.ref is not None

    data_row = header_row + 1
    row_values = [ws.cell(data_row, c + 1).value for c in range(len(DEFAULT_COLUMNS))]
    assert "Robert Niyitanga" in row_values


def test_xlsx_includes_meta_header_fields():
    _, content, _ = build_users_export(SAMPLE_ROWS, DEFAULT_COLUMNS, "xlsx", "superadmin@pesaa.io", {"role": "admin"})
    wb = load_workbook(io.BytesIO(content))
    ws = wb.active
    all_text = " ".join(str(ws.cell(r, 1).value) for r in range(1, 12) if ws.cell(r, 1).value)
    assert "PESAA" in all_text
    assert "superadmin@pesaa.io" in all_text
    assert "Total Users: 2" in all_text
    assert "Role: admin" in all_text


def test_xlsx_empty_result_still_produces_valid_file():
    _, content, _ = build_users_export([], DEFAULT_COLUMNS, "xlsx", "a@pesaa.io", {})
    wb = load_workbook(io.BytesIO(content))
    assert wb.active is not None
